import os
import sys
import argparse
from pathlib import Path
from io import BytesIO

# Office 文档内部是 XML 结构。下面这些命名空间和标签名用于直接读取
# Word 段落、表格、文字 run、样式和图片引用，避免只依赖 python-docx 的高层 API。
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
V_NS = 'urn:schemas-microsoft-com:vml'

TAG_P = f'{{{W_NS}}}p'
TAG_R = f'{{{W_NS}}}r'
TAG_T = f'{{{W_NS}}}t'
TAG_PPR = f'{{{W_NS}}}pPr'
TAG_PSTYLE = f'{{{W_NS}}}pStyle'
TAG_TR = f'{{{W_NS}}}tr'
TAG_TC = f'{{{W_NS}}}tc'
TAG_BLIP = f'{{{A_NS}}}blip'
TAG_IMAGEDATA = f'{{{V_NS}}}imagedata'


def _save_image(image_data, ext, image_dir, counter, min_size=20):
    """保存从文档中提取出的图片，并返回图片路径。

    通过 Pillow 检查图片尺寸，跳过宽或高小于 min_size 像素的图片。
    Word 文档中经常嵌入极小的占位图、装饰元素或格式化辅助图片，
    这些图片对内容没有意义，保存只会污染图片目录。
    """
    if ext == 'jpeg':
        ext = 'jpg'

    # Word 里可能包含 wmf/emf 这类矢量图。Markdown 浏览器对它们支持不好，
    # 这里尽量转成 PNG；如果转换失败，就跳过这张图片。
    if ext in ('x-wmf', 'wmf', 'emf', 'x-emf'):
        try:
            from PIL import Image
            img = Image.open(BytesIO(image_data))
            w, h = img.size
            if w < min_size or h < min_size:
                return None
            img = img.convert('RGB')
            buf = BytesIO()
            img.save(buf, format='PNG')
            image_data = buf.getvalue()
            ext = 'png'
        except Exception:
            return None
    else:
        # 非矢量图也需要检查尺寸，过滤掉过小的占位图和装饰元素。
        try:
            from PIL import Image
            img = Image.open(BytesIO(image_data))
            w, h = img.size
            if w < min_size or h < min_size:
                return None
        except Exception:
            # 无法读取尺寸时仍然保存，避免丢失有效图片。
            pass

    counter[0] += 1
    image_name = f"image_{counter[0]}.{ext}"
    image_path = image_dir / image_name
    with open(image_path, 'wb') as f:
        f.write(image_data)
    return image_path


def _extract_run_images(run_elem, doc, image_dir, counter):
    """从 Word 的一个 run 中提取图片。

    Word 图片有两种常见存放形式：
    - DrawingML: 现代 docx 常见格式，对应 blip 标签
    - VML: 旧版兼容格式，对应 imagedata 标签
    """
    results = []

    # 处理现代 docx 图片引用。
    for blip in run_elem.findall(f'.//{TAG_BLIP}'):
        rId = blip.get(f'{{{R_NS}}}embed')
        if rId and rId in doc.part.rels:
            try:
                rel = doc.part.rels[rId]
                ext = rel.target_part.content_type.split('/')[-1]
                img_path = _save_image(rel.target_part.blob, ext, image_dir, counter)
                if img_path:
                    results.append(img_path)
            except Exception:
                pass

    # 处理旧版 VML 图片引用。
    for imagedata in run_elem.findall(f'.//{TAG_IMAGEDATA}'):
        rId = imagedata.get(f'{{{R_NS}}}id')
        if rId and rId in doc.part.rels:
            try:
                rel = doc.part.rels[rId]
                ext = rel.target_part.content_type.split('/')[-1]
                img_path = _save_image(rel.target_part.blob, ext, image_dir, counter)
                if img_path:
                    results.append(img_path)
            except Exception:
                pass
    return results


def _get_style(para_elem):
    """读取 Word 段落样式，用于判断标题、列表等结构。"""
    pPr = para_elem.find(TAG_PPR)
    if pPr is not None:
        pStyle = pPr.find(TAG_PSTYLE)
        if pStyle is not None:
            return pStyle.get(f'{{{W_NS}}}val', '')
    return ''


def _format_text(text, style):
    """把 Word 样式转换成 Markdown 写法。"""
    if style.startswith('Heading'):
        try:
            level = int(style.replace('Heading ', ''))
        except ValueError:
            level = 1
        return f"{'#' * level} {text}"
    if style == 'List Bullet':
        return f"- {text}"
    if style == 'List Number':
        return f"1. {text}"
    return text


def convert_word_to_md(input_path, output_path, image_dir=None):
    """将 Word docx 转为 Markdown。

    当前保留三类主要信息：段落文本、表格、内嵌图片。
    图片会写入 output_path 同级的 images 目录，并在 Markdown 中使用相对路径引用。
    """
    try:
        from docx import Document
    except ImportError:
        print("错误: 请安装python-docx库: pip install python-docx")
        return False

    try:
        doc = Document(input_path)
        md_content = []
        counter = [0]

        if image_dir is None:
            image_dir = Path(output_path).parent / "images"
        image_dir = Path(image_dir)
        image_dir.mkdir(parents=True, exist_ok=True)

        out_parent = Path(output_path).parent

        # 逐个读取 Word 正文中的顶层元素，尽量保持原文顺序。
        for element in doc.element.body:
            local_tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag

            if local_tag == 'p':
                style = _get_style(element)
                text_buf = []
                has_content = False

                # 一个段落可能由多个 run 组成；图片也经常挂在 run 里面。
                for run in element.findall(TAG_R):
                    images = _extract_run_images(run, doc, image_dir, counter)
                    if images:
                        # 如果图片前面已经积累了文字，先把文字写入 Markdown，
                        # 再插入图片，避免文字和图片顺序错乱。
                        accumulated = ''.join(text_buf).strip()
                        if accumulated:
                            md_content.append(_format_text(accumulated, style))
                            text_buf = []
                            has_content = True
                        for img_path in images:
                            rel_path = os.path.relpath(img_path, out_parent)
                            md_content.append(f"![图片]({rel_path})")
                            has_content = True
                    else:
                        # 普通文字 run：收集所有文本节点，最后合并成一个段落。
                        for t in run.findall(TAG_T):
                            if t.text:
                                text_buf.append(t.text)

                remaining = ''.join(text_buf).strip()
                if remaining:
                    md_content.append(_format_text(remaining, style))
                    has_content = True

                if has_content:
                    md_content.append("")

            elif local_tag == 'tbl':
                # Word 表格转成 Markdown 表格。第一行默认作为表头。
                rows = element.findall(f'.//{TAG_TR}')
                for row_idx, row in enumerate(rows):
                    cells = row.findall(TAG_TC)
                    cell_texts = []
                    for cell in cells:
                        parts = []
                        for t in cell.findall(f'.//{TAG_T}'):
                            if t.text:
                                parts.append(t.text)
                        cell_texts.append(''.join(parts).strip())
                    md_content.append('| ' + ' | '.join(cell_texts) + ' |')
                    if row_idx == 0:
                        md_content.append('| ' + ' | '.join(['---'] * len(cell_texts)) + ' |')
                md_content.append("")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(md_content))
        return True
    except Exception as e:
        print(f"转换Word文档失败: {e}")
        import traceback
        traceback.print_exc()
        return False

def convert_ppt_to_md(input_path, output_path, image_dir=None):
    """将 PPT/PPTX 转为 Markdown。

    每一页幻灯片会转换成一个二级标题，页面中的文本框和图片按读取顺序写入。
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from PIL import Image
    except ImportError:
        print("错误: 请安装python-pptx和Pillow库: pip install python-pptx Pillow")
        return False

    try:
        prs = Presentation(input_path)
        md_content = []
        if image_dir is None:
            image_dir = Path(output_path).parent / "images"
        image_dir = Path(image_dir)
        image_dir.mkdir(parents=True, exist_ok=True)

        # 幻灯片天然按页组织，用“## 幻灯片 N”保留页结构。
        for slide_num, slide in enumerate(prs.slides, 1):
            md_content.append(f"## 幻灯片 {slide_num}")
            md_content.append("")

            for shape in slide.shapes:
                # shape_type == 13 表示图片。图片不走文本分支，避免重复处理。
                if hasattr(shape, "text") and shape.text.strip():
                    if shape.shape_type == 13:
                        pass
                    else:
                        md_content.append(shape.text)
                        md_content.append("")

                if shape.shape_type == 13:
                    # 提取 PPT 中的原始图片二进制，保存后写入 Markdown 图片链接。
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        ext = image.ext
                        if ext == 'jpeg':
                            ext = 'jpg'
                        image_name = f"slide{slide_num}_image{len(list(image_dir.glob('*')))}.{ext}"
                        image_path = image_dir / image_name

                        with open(image_path, 'wb') as f:
                            f.write(image_bytes)

                        rel_path = os.path.relpath(image_path, Path(output_path).parent)
                        md_content.append(f"![幻灯片{slide_num}图片]({rel_path})")
                        md_content.append("")
                    except Exception as e:
                        print(f"警告: 无法处理幻灯片{slide_num}中的图片: {e}")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(md_content))
        return True
    except Exception as e:
        print(f"转换PPT文档失败: {e}")
        return False

def convert_pdf_to_md(input_path, output_path):
    """将 PDF 转为 Markdown。

    PDF 的结构通常比 Word/PPT 更弱，这里按页提取文本和页面中的图片区域。
    扫描版 PDF 可能没有可提取文本，需要后续接 OCR 才能获得正文。
    """
    try:
        import pdfplumber
        from PIL import Image
    except ImportError:
        print("错误: 请安装pdfplumber和Pillow库: pip install pdfplumber Pillow")
        return False

    try:
        md_content = []
        image_dir = Path(output_path).parent / "images"
        image_dir.mkdir(exist_ok=True)

        with pdfplumber.open(input_path) as pdf:
            # 按页处理，方便后续追溯内容来自 PDF 的哪一页。
            for page_num, page in enumerate(pdf.pages, 1):
                md_content.append(f"## 第 {page_num} 页")
                md_content.append("")

                # extract_text 只适用于包含文本层的 PDF。
                text = page.extract_text()
                if text:
                    md_content.append(text)
                    md_content.append("")

                # 提取 PDF 页面中识别到的图片区域，并以 PNG 保存。
                for img_index, img in enumerate(page.images):
                    try:
                        x0 = img['x0']
                        y0 = img['top']
                        x1 = img['x1']
                        y1 = img['bottom']

                        cropped = page.crop((x0, y0, x1, y1))
                        img_obj = cropped.to_image()
                        img_bytes = BytesIO()
                        img_obj.save(img_bytes, format='PNG')
                        img_bytes.seek(0)

                        image_name = f"page{page_num}_img{img_index}.png"
                        image_path = image_dir / image_name

                        with open(image_path, 'wb') as f:
                            f.write(img_bytes.getvalue())

                        rel_path = os.path.relpath(image_path, Path(output_path).parent)
                        md_content.append(f"![第{page_num}页图片]({rel_path})")
                        md_content.append("")
                    except Exception as e:
                        print(f"警告: 无法处理第{page_num}页的图片: {e}")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(md_content))
        return True
    except Exception as e:
        print(f"转换PDF文档失败: {e}")
        return False

def main():
    """命令行入口：转换单个 Word/PPT/PDF 文件。"""
    parser = argparse.ArgumentParser(description='将Word、PPT、PDF文档转换为Markdown格式')
    parser.add_argument('input', help='输入文件路径')
    parser.add_argument('-o', '--output', help='输出文件路径（可选）')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"错误: 输入文件不存在: {input_path}")
        sys.exit(1)

    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_suffix('.md')

    output_path.parent.mkdir(parents=True, exist_ok=True)

    ext = input_path.suffix.lower()
    success = False

    # 根据扩展名选择对应的转换函数。
    if ext == '.docx':
        print(f"正在转换Word文档: {input_path}")
        success = convert_word_to_md(input_path, output_path)
    elif ext == '.pptx':
        print(f"正在转换PPT文档: {input_path}")
        success = convert_ppt_to_md(input_path, output_path)
    elif ext == '.pdf':
        print(f"正在转换PDF文档: {input_path}")
        success = convert_pdf_to_md(input_path, output_path)
    else:
        print(f"错误: 不支持的文件格式: {ext}")
        print("支持的格式: .docx, .pptx, .pdf")
        sys.exit(1)

    if success:
        print(f"转换成功: {output_path}")
        sys.exit(0)
    else:
        print("转换失败")
        sys.exit(1)

if __name__ == "__main__":
    main()
